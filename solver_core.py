# solver_core.py
from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from ortools.sat.python import cp_model

TIME_LIMIT_SECONDS = 60
NEXT_DAY_MIN_AVAILABLE = 5
PASS, FAIL, PENDING = "通過", "違規", "待確認"

def as_date(value):
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if isinstance(value, str):
        try: return date.fromisoformat(value.strip())
        except ValueError: return None
    return None

def text_value(value):
    return "" if value is None else str(value).strip()

def is_night_text(s):
    return "大夜" in s or "夜班" in s or "20:30" in s

def is_late_text(s):
    return "晚班" in s or "15:00" in s or s == "晚"

def shift_workload(shifts):
    result = 0
    for shift in shifts:
        s = text_value(shift)
        if is_night_text(s): result = max(result, 2)
        elif "早班" in s or is_late_text(s) or s == "早" or "09:00" in s: result = max(result, 1)
    return result

# （此處保留你原本的 validate 函數...）
# 為了節省篇幅，validate 內容與你原本的一致，確保它完整存在於 solver_core.py 中

def run_solver(file_path: Path):
    """供網頁呼叫的主執行函數"""
    wb = load_workbook(file_path)
    # ... (將你原本 main() 裡面的讀取、模型建立、求解與寫入 Excel 邏輯全部貼在這邊)
    # 只要將原本的 FILE 改成傳入的 file_path 即可
    
    # 範例回傳執行結果或訊息
    return True, "排班運算成功完成！"